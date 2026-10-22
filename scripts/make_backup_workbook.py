"""Build the Week 11 group backup workbook (plain Excel, one file per group).

Run:  python3 scripts/make_backup_workbook.py
Writes public/backup/Week11-group-backup.xlsx. Yellow cells are for students, grey italic rows are examples.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "public" / "backup" / "Week11-group-backup.xlsx"

FONT = "Arial"
INK = "2B2622"
ACCENT = "7D5236"
TINT = "F3EADD"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor=ACCENT)
TINT_FILL = PatternFill("solid", fgColor=TINT)
EX_FILL = PatternFill("solid", fgColor="EFEFEF")
THIN = Side(style="thin", color="BFB6AB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

f_body = Font(name=FONT, size=10, color=INK)
f_bold = Font(name=FONT, size=10, bold=True, color=INK)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True, color=ACCENT)
f_sub = Font(name=FONT, size=10, italic=True, color="6F665E")
f_ex = Font(name=FONT, size=10, italic=True, color="6F665E")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

wb = Workbook()


def sheet(name, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f_title
    ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f_sub
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ws.max_column, 4))
        ws.row_dimensions[2].height = 30


def header(ws, row, labels):
    for c, text in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.font = f_head
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = 32


def put(ws, row, col, value, style="body", fill=None, height=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = {"body": f_body, "bold": f_bold, "ex": f_ex, "sub": f_sub}[style]
    cell.alignment = WRAP
    cell.border = BORDER
    if fill:
        cell.fill = fill
    if height:
        ws.row_dimensions[row].height = height
    return cell


def input_rows(ws, first_row, count, ncols, height=42, start_col=1):
    for r in range(first_row, first_row + count):
        for c in range(start_col, start_col + ncols):
            put(ws, r, c, None, fill=INPUT_FILL)
        ws.row_dimensions[r].height = height


def example_row(ws, row, values, height=42):
    for c, v in enumerate(values, start=1):
        put(ws, row, c, v, style="ex", fill=EX_FILL)
    ws.row_dimensions[row].height = height


def dropdown(ws, cell_range, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = "Pick one from the list"
    ws.add_data_validation(dv)
    dv.add(cell_range)


def field_table(ws, first_row, rows, ans_height=70):
    """Rows of (label, example). Column A label, B answer (yellow), C example (grey)."""
    header(ws, first_row, ["What to fill in", "Your group's answer", "Example (delete nothing, just write in yellow)"])
    for i, (label, example) in enumerate(rows, start=first_row + 1):
        put(ws, i, 1, label, style="bold", fill=TINT_FILL)
        put(ws, i, 2, None, fill=INPUT_FILL)
        put(ws, i, 3, example, style="ex", fill=EX_FILL)
        ws.row_dimensions[i].height = ans_height


# ------------------------------------------------------------------ Start here
ws = sheet("Start here", [34, 26, 26, 40], first=True)
title(ws, "Week 11: From Theme to Implication (group backup)", "A plain backup for your group. Use it if the website is slow or a device is unavailable. One person types for the group; share the file so everyone can see it.")
put(ws, 4, 1, "Group name or number", style="bold", fill=TINT_FILL)
put(ws, 4, 2, None, fill=INPUT_FILL)
ws.merge_cells("B4:D4")
ws.row_dimensions[4].height = 24
put(ws, 6, 1, "Who is in the group?", style="bold", fill=TINT_FILL)
header(ws, 7, ["Initials (no full names needed)", "Video you watched", "Notes", ""])
ws.merge_cells("C7:D7")
for r in range(8, 14):
    put(ws, r, 1, None, fill=INPUT_FILL)
    put(ws, r, 2, None, fill=INPUT_FILL)
    put(ws, r, 3, None, fill=INPUT_FILL)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 22
dropdown(ws, "B8:B13", ["Simon", "Sarah", "Stacey"])
put(ws, 15, 1, "Research question", style="bold", fill=TINT_FILL)
put(ws, 15, 2, "How do individuals experience and make sense of their relationships with technology in everyday life?", fill=None)
ws.merge_cells("B15:D15")
ws.row_dimensions[15].height = 34
put(ws, 17, 1, "Tabs in this workbook", style="bold", fill=TINT_FILL)
tabs = [
    ("T1 Codes", "Task 1 Step 1: everyone's codes (10 min, individual)"),
    ("T1 Cluster", "Task 1 Step 2: share, cluster, discussion notes (8 min, group)"),
    ("T1 Theme", "Task 1 Step 3: your candidate theme (7 min, group)"),
    ("T2 Critique", "Task 2 Step 1: critique Theme A and Theme B (10 min, group)"),
    ("T2 Refine", "Task 2 Steps 2 and 3: refine the weaker theme, revisit yours (15 min, group)"),
    ("T3 Implication", "Task 3: write and share one implication (25 min)"),
    ("T4 Ed post", "Task 4: draft your reflection post (10 min, individual)"),
]
for i, (t, d) in enumerate(tabs, start=18):
    put(ws, i, 1, t, style="bold")
    put(ws, i, 2, d)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
put(ws, 26, 1, "Colour key", style="bold", fill=TINT_FILL)
put(ws, 26, 2, "Yellow = fill in", fill=INPUT_FILL)
put(ws, 26, 3, "Grey italic = example", style="ex", fill=EX_FILL)
put(ws, 26, 4, "White = fixed text")
ws.row_dimensions[26].height = 22
put(ws, 28, 1, "Privacy", style="bold", fill=TINT_FILL)
put(ws, 28, 2, "Use initials rather than full names, and keep personal details out of the cells.")
ws.merge_cells("B28:D28")

# ------------------------------------------------------------------ T1 Codes
ws = sheet("T1 Codes", [10, 12, 8, 46, 24, 40])
title(ws, "Task 1, Step 1: initial codes", "Each person codes 3-5 moments from their own video. O = you saw or heard it, R = the participant reports it, I = your interpretation. A code is a short label for what the moment is about.")
header(ws, 4, ["Who", "Video", "O / R / I", "Observation or excerpt (with timestamp)", "Initial code", "Why did you assign this code?"])
example_row(ws, 5, ["Example", "Simon", "R", "Simon 3:52, says without it he would be totally dependent on other people", "Independence through technology", "He links the equipment directly to doing things when he wants to."], 48)
input_rows(ws, 6, 25, 6, 44)
dropdown(ws, "B6:B30", ["Simon", "Sarah", "Stacey"])
dropdown(ws, "C6:C30", ["O", "R", "I"])
ws.freeze_panes = "A5"

# ------------------------------------------------------------------ T1 Cluster
ws = sheet("T1 Cluster", [38, 42, 44])
title(ws, "Task 1, Step 2: share and cluster", "Take turns introducing your video and sharing 2-3 observations with their codes. Group codes that belong together, then say what holds each cluster together.")
header(ws, 4, ["Related observations and codes", "What connects or distinguishes them?", "What might this suggest about people's relationships with technology?"])
example_row(ws, 5, ["Simon coffee request (6:13); Sarah lets Sinclair decide (9:28)", "Both choose who or what they lean on.", "People shape their own dependence."], 48)
input_rows(ws, 6, 8, 3, 52)
header(ws, 16, ["Discussion prompt", "Notes from your discussion", ""])
ws.merge_cells("B16:C16")
prompts = [
    "What similarities and differences do you notice across the experiences in the different videos?",
    "What underlying meanings might connect your observations?",
    "Why might group members have interpreted similar experiences differently?",
    "Which observations challenge the interpretation that is starting to emerge?",
]
for i, p in enumerate(prompts, start=17):
    put(ws, i, 1, p, style="bold", fill=TINT_FILL)
    put(ws, i, 2, None, fill=INPUT_FILL)
    put(ws, i, 3, None, fill=INPUT_FILL)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
    ws.row_dimensions[i].height = 58

# ------------------------------------------------------------------ T1 Theme
ws = sheet("T1 Theme", [34, 54, 46])
title(ws, "Task 1, Step 3: candidate theme", "A theme is a pattern of shared meaning organised around a central concept. It interprets what people said rather than summarising it. Keep this theme for Tasks 2 and 3.")
field_table(ws, 4, [
    ("Working name for the theme", "People choose what they depend on"),
    ("Central organising concept", "Chosen dependence: agency shown through who or what people rely on"),
    ("2-3 supporting observations (ideally from different videos, with timestamps)", "Simon 6:13 asks carers for coffee; Sarah 9:36 calls it an agreed dynamic; Stacey 2:10 chooses Aibo"),
    ("What does this theme help us understand about relationships with technology?", "Reliance is something people actively arrange, not only something that happens to them"),
])

# ------------------------------------------------------------------ T2 Critique
ws = sheet("T2 Critique", [40, 40, 40, 30])
title(ws, "Task 2, Step 1: critique the two themes", "Judge both themes. Quote timestamps when you cite evidence, and say how else the evidence could be read.")
put(ws, 4, 1, "THEME A", style="bold", fill=TINT_FILL)
put(ws, 4, 2, "Technology makes everyday life better for people.", style="bold")
ws.merge_cells("B4:D4")
ws.row_dimensions[4].height = 24
a_ev = [
    "Simon controls his TV, computer and wheelchair with head-operated switches (R, 1:38 to 2:31).",
    "Sinclair monitors Sarah's heart rate and reminds her to drink water (R, 4:00).",
    "Mia comes over on her own asking for attention and pets (R, 5:06).",
]
for i, t in enumerate(a_ev, start=5):
    put(ws, i, 1, "Evidence" if i == 5 else None, style="bold")
    put(ws, i, 2, t)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    ws.row_dimensions[i].height = 30
put(ws, 9, 1, "THEME B", style="bold", fill=TINT_FILL)
put(ws, 9, 2, "People use technology to choose who and what they depend on, and they value that arrangement because they chose it, even when others see it as a problem.", style="bold")
ws.merge_cells("B9:D9")
ws.row_dimensions[9].height = 44
b_ev = [
    'Simon: "without it I would be totally dependent on other people" (R, 3:52), yet still asks his carers to make him a coffee (R, 6:13); email lets him keep secrets from his parents (R, 8:37).',
    'Sarah: Sinclair "takes that weight ... makes the decisions" (R, 9:28); "an agreed-upon dynamic" (R, 9:36), in reply to the presenter calling him "quite controlling" (9:03).',
    "Stacey: after months of round-the-clock care for her dying dog (R, 1:49), she was not ready for another dog and chose Aibo (R, 2:10).",
]
for i, t in enumerate(b_ev, start=10):
    put(ws, i, 1, "Evidence" if i == 10 else None, style="bold")
    put(ws, i, 2, t)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    ws.row_dimensions[i].height = 44
header(ws, 14, ["Evaluation criteria", "Theme A", "Theme B", ""])
ws.merge_cells("C14:D14")
for i, crit in enumerate(["Central organising concept", "Supporting observations", "Strengths", "Limitations or alternative interpretations"], start=15):
    put(ws, i, 1, crit, style="bold", fill=TINT_FILL)
    put(ws, i, 2, None, fill=INPUT_FILL)
    put(ws, i, 3, None, fill=INPUT_FILL)
    ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=4)
    ws.row_dimensions[i].height = 62
put(ws, 19, 1, "How might your own assumptions affect which theme you find convincing?", style="bold", fill=TINT_FILL)
put(ws, 19, 2, None, fill=INPUT_FILL)
ws.merge_cells("B19:D19")
ws.row_dimensions[19].height = 62

header(ws, 21, ["Quick rating: does the theme meet this?", "Theme A", "Theme B", "Our theme (Task 1)"])
crit6 = [
    "Has a clear theme name",
    "Has a central organising concept",
    "Has evidence from 2+ interviews",
    "Goes beyond description",
    "Engages with contradiction",
    "Answers the research question",
]
for i, c in enumerate(crit6, start=22):
    put(ws, i, 1, c, style="bold", fill=TINT_FILL)
    for col in (2, 3, 4):
        cell = put(ws, i, col, None, fill=INPUT_FILL)
        cell.alignment = CENTER
    ws.row_dimensions[i].height = 22
dropdown(ws, "B22:D27", ["Yes", "Partly", "No"])
put(ws, 28, 1, "Criteria met (count of Yes)", style="bold")
for col, letter in ((2, "B"), (3, "C"), (4, "D")):
    cell = put(ws, 28, col, f'=COUNTIF({letter}22:{letter}27,"Yes")', style="bold")
    cell.alignment = CENTER

# ------------------------------------------------------------------ T2 Refine
ws = sheet("T2 Refine", [34, 54, 46])
title(ws, "Task 2, Steps 2 and 3: refine and revisit", "Revise the weaker theme (name, central concept, evidence or interpretation), test it, then go back to your own Task 1 theme.")
field_table(ws, 4, [
    ("Which theme did your group judge weaker?", "Theme A"),
    ("Revised theme name", "Technology as an extension of what people want to do"),
    ("Central organising concept", "Technology widens agency, with limits that people notice"),
    ("Evidence (with timestamps)", "Simon 3:52 and 6:13; Stacey 3:20 to 3:50 when Mia is grumpy after being off"),
    ("Interpretation: what does the theme say is going on?", "People value technology when it serves their own aims, and they notice where it falls short"),
], ans_height=62)
header(ws, 11, ["Test your revision", "Your answer", ""])
ws.merge_cells("B11:C11")
tests = [
    "Does your revised theme explain an underlying pattern, rather than only naming a topic?",
    "What connects the observations across the interviews?",
    "Does it help answer the research question?",
    "Do any differences or contradictions complicate it?",
    "How might another researcher interpret the same observations?",
]
for i, q in enumerate(tests, start=12):
    put(ws, i, 1, q, style="bold", fill=TINT_FILL)
    put(ws, i, 2, None, fill=INPUT_FILL)
    put(ws, i, 3, None, fill=INPUT_FILL)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
    ws.row_dimensions[i].height = 52
dropdown(ws, "B5", ["Theme A", "Theme B"])
field_table(ws, 18, [
    ("One change that would strengthen your original theme (Task 1)", "Add a contradicting observation from another video"),
    ("Why would that change help?", "It shows the theme holds up when it is challenged"),
    ("Has your understanding of the interviews, or of your own assumptions, changed?", "We had assumed technology use was mostly positive"),
], ans_height=56)

# ------------------------------------------------------------------ T3 Implication
ws = sheet("T3 Implication", [34, 60, 46])
title(ws, "Task 3: from theme to implication", "A theme says what is going on. An implication names who should do or think differently, and what. Write one in 2-3 sentences that you could not have written without the data.")
put(ws, 4, 1, "Weak example", style="bold", fill=TINT_FILL)
put(ws, 4, 2, "Designers should make AI study tools more trustworthy. (You could write this without any data.)", style="ex", fill=EX_FILL)
put(ws, 5, 1, "Stronger example", style="bold", fill=TINT_FILL)
put(ws, 5, 2, "AI study tools could show sources and confidence when a student prepares for an assessment, rather than in every interaction.", style="ex", fill=EX_FILL)
ws.merge_cells("B4:C4")
ws.merge_cells("B5:C5")
ws.row_dimensions[4].height = 32
ws.row_dimensions[5].height = 32
put(ws, 7, 1, "Our theme (from Task 1, or the revised one)", style="bold", fill=TINT_FILL)
put(ws, 7, 2, None, fill=INPUT_FILL)
ws.merge_cells("B7:C7")
ws.row_dimensions[7].height = 48
put(ws, 8, 1, "Our implication (2-3 sentences)", style="bold", fill=TINT_FILL)
put(ws, 8, 2, None, fill=INPUT_FILL)
ws.merge_cells("B8:C8")
ws.row_dimensions[8].height = 110
put(ws, 9, 1, "Sentences (aim for 2-3)", style="bold")
count = '=(LEN(B8)-LEN(SUBSTITUTE(B8,".","")))+(LEN(B8)-LEN(SUBSTITUTE(B8,"?","")))+(LEN(B8)-LEN(SUBSTITUTE(B8,"!","")))'
put(ws, 9, 2, count)
ws.merge_cells("B9:C9")
ws.cell(row=9, column=2).alignment = Alignment(horizontal="left")
checks = [
    "Does it say something the theme does not already say (not a restatement)?",
    "Could we NOT have written it without this data?",
]
header(ws, 11, ["Self-check", "Yes or No", "Notes"])
for i, c in enumerate(checks, start=12):
    put(ws, i, 1, c, style="bold", fill=TINT_FILL, height=36)
    cell = put(ws, i, 2, None, fill=INPUT_FILL)
    cell.alignment = CENTER
    put(ws, i, 3, None, fill=INPUT_FILL)
dropdown(ws, "B12:B13", ["Yes", "No"])
put(ws, 15, 1, "For the class share (one line each)", style="bold", fill=TINT_FILL)
put(ws, 15, 2, None, fill=INPUT_FILL)
put(ws, 16, 1, "Theme on our card", style="bold")
put(ws, 16, 2, None, fill=INPUT_FILL)
put(ws, 17, 1, "Implication on our card", style="bold")
put(ws, 17, 2, None, fill=INPUT_FILL)
ws.merge_cells("B16:C16")
ws.merge_cells("B17:C17")
ws.merge_cells("B15:C15")
ws["A15"].value = "Class share card"
ws["B15"].value = "Copy your theme and implication into the two rows below so the tutor can project them."
ws["B15"].fill = PatternFill(fill_type=None)
ws["B15"].font = f_sub
ws.row_dimensions[16].height = 40
ws.row_dimensions[17].height = 70

# ------------------------------------------------------------------ T4 Ed post
ws = sheet("T4 Ed post", [12, 70, 12, 16])
title(ws, "Task 4: reflection post (individual)", "Answer both prompts in about 100-150 words using examples from your own analysis: (1) Why might researchers develop different themes from the same qualitative dataset, and why does reflexivity matter? (2) What did you have to add to get from your theme to your implication? Paste the final text into Ed yourself.")
ws.row_dimensions[2].height = 58
header(ws, 4, ["Who", "Draft", "Words", "Status"])
example_row(ws, 5, ["Example", "Different researchers bring different assumptions, so the same moment can be coded differently. For example, I coded Simon's coffee request as independence because I assumed technology mostly helps, while a groupmate coded the same moment as dependence because they noticed the carers. Neither reading is wrong, but each shows where we were coming from. Reflexivity matters because it asks researchers to say where they stand, so readers can judge how that position shaped the themes. To get from our theme to our implication I had to add a who and a what: who should act differently, and what they should do.", None, None], 76)
ws["C5"] = '=IF(TRIM(B5)="",0,LEN(TRIM(B5))-LEN(SUBSTITUTE(TRIM(B5)," ",""))+1)'
ws["D5"] = '=IF(C5<100,"Under 100",IF(C5>150,"Over 150","In range"))'
for r in range(6, 11):
    put(ws, r, 1, None, fill=INPUT_FILL)
    put(ws, r, 2, None, fill=INPUT_FILL)
    c = put(ws, r, 3, f'=IF(TRIM(B{r})="",0,LEN(TRIM(B{r}))-LEN(SUBSTITUTE(TRIM(B{r})," ",""))+1)', style="bold")
    c.alignment = CENTER
    d = put(ws, r, 4, f'=IF(C{r}<100,"Under 100",IF(C{r}>150,"Over 150","In range"))', style="bold")
    d.alignment = CENTER
    ws.row_dimensions[r].height = 150
for c in ("C5", "D5"):
    ws[c].font = f_ex
    ws[c].fill = EX_FILL
    ws[c].border = BORDER
    ws[c].alignment = CENTER

wb.calculation.fullCalcOnLoad = True  # no cached values are stored; Excel and Sheets compute the counts on open
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("wrote", OUT)
